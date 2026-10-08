"""独立校验: 不复用求解器的中间结果, 直接读输出的 Excel 反查每条约束.

目的是确认「求解器说它满足约束」和「结果真的满足约束」是两件独立验证过的事.
"""

import sys
import openpyxl
import config
import commute
from metro import TravelTimeEngine
from parse_input import load

HERE = __file__.rsplit('/', 1)[0]


def data_rows(ws, first_header):
    """定位表头行后返回数据行, 不假设表头固定在第几行."""
    hdr = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=4, values_only=True), 1):
        if row and str(row[0]).strip() == first_header:
            hdr = i
            break
    if hdr is None:
        raise RuntimeError(f'{ws.title}: 找不到表头 {first_header!r}')
    return [r for r in ws.iter_rows(min_row=hdr + 1, values_only=True) if r[0] not in (None, '')]


def main(src, out):
    eng = TravelTimeEngine(f'{HERE}/sh_metro_raw.json')
    vols, classes, sites, _ = load(src, eng)
    prov, _ = commute.build(eng)
    V = {v['id']: v for v in vols}
    C = {(c['site'], c['cls']): c for c in classes}

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb['分配理由']
    rows = []
    for r in data_rows(ws, '项目点'):
        rows.append({'site': str(r[0]), 'cls': str(r[1]), 'slot': r[2], 'vol': str(r[3]),
                     'minutes': float(r[4]), 'transfers': int(r[5]),
                     'from': r[6], 'to': r[7]})

    fails, checks = [], 0

    # 1 时间硬约束
    for a in rows:
        checks += 1
        c = C.get((a['site'], a['cls']))
        if c is None:
            fails.append(f"{a['vol']}: 班级 {a['site']}-{a['cls']} 不存在")
            continue
        if c['slot'] not in V[a['vol']]['slots']:
            fails.append(f"{a['vol']} 被排到 {c['slot']}, 但其可用时段是 "
                         f"{V[a['vol']]['slots']}")
        if a['slot'] != c['slot']:
            fails.append(f"{a['vol']}: 输出时段 {a['slot']} 与班级时段 {c['slot']} 不一致")

    # 2 每人最多一个班
    cnt = {}
    for a in rows:
        cnt[a['vol']] = cnt.get(a['vol'], 0) + 1
    for v, n in cnt.items():
        checks += 1
        if n > config.MAX_CLASS_PER_VOLUNTEER:
            fails.append(f'{v} 被分配了 {n} 个班, 超过上限 {config.MAX_CLASS_PER_VOLUNTEER}')

    # 3 班级人数不超需求
    per = {}
    for a in rows:
        per[(a['site'], a['cls'])] = per.get((a['site'], a['cls']), 0) + 1
    for k, n in per.items():
        checks += 1
        if n > C[k]['need']:
            fails.append(f'{k} 分到 {n} 人, 超过需求 {C[k]["need"]}')

    # 4 通勤时间重算一遍而不是信输出值
    for a in rows:
        checks += 1
        v = V[a['vol']]
        r = prov.pair(v['stations'], C[(a['site'], a['cls'])]['stations'])
        if r is None:
            fails.append(f"{a['vol']} -> {a['site']}-{a['cls']}: 重算不出通勤时间")
            continue
        if abs(r[0] - a['minutes']) > 0.15:
            fails.append(f"{a['vol']} -> {a['site']}-{a['cls']}: 输出 {a['minutes']} 分钟, "
                         f"重算 {r[0]:.1f} 分钟")
        if a['from'] not in v['stations']:
            fails.append(f"{a['vol']}: 起点站 {a['from']} 不在其填报站点 {v['stations']} 内")
        if a['to'] not in C[(a['site'], a['cls'])]['stations']:
            fails.append(f"{a['site']}-{a['cls']}: 终点站 {a['to']} 不在项目点站点内")

    # 5 缺口表与分配表加起来等于需求
    short = sum(int(r[6]) for r in data_rows(wb['未填补名额'], '项目点'))
    checks += 1
    need = sum(c['need'] for c in classes)
    if len(rows) + short != need:
        fails.append(f'名额不闭合: 已分配 {len(rows)} + 缺口 {short} != 需求 {need}')

    # 6 未安置志愿者表与分配表无交集且总数对得上
    unp = [str(r[0]) for r in data_rows(wb['未安置志愿者'], '志愿者')]
    checks += 1
    if set(unp) & set(cnt):
        fails.append(f'同时出现在已分配和未安置名单: {set(unp) & set(cnt)}')
    checks += 1
    if len(unp) + len(cnt) != len(vols):
        fails.append(f'人数不闭合: 已安置 {len(cnt)} + 未安置 {len(unp)} != 总数 {len(vols)}')

    print(f'共检查 {checks} 项')
    if fails:
        print(f'发现 {len(fails)} 处问题:')
        for f in fails:
            print('  ✗ ' + f)
        return 1
    print('✓ 全部通过: 时间匹配、每人一班、班级人数上限、通勤时间、'
          '起终点站归属、名额闭合、人数闭合')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/input_25spring.xlsx',
                  sys.argv[2] if len(sys.argv) > 2 else f'{HERE}/output_分配结果.xlsx'))
