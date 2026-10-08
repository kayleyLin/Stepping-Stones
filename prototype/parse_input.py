"""数据清洗层: 把原始 SET Allocation Survey 解析成结构化记录.

这一层的存在本身就是《数据需求规格书》的论据. 每一条 _warn 都对应
一处本可以在采集阶段用下拉选项或固定分隔符消除的格式问题.
"""

import re
import openpyxl
from metro import norm_station

# ---- 时段归一化 ----
DAY = {'monday': '周一', 'tuesday': '周二', 'wednesday': '周三', 'thursday': '周四',
       'friday': '周五', 'saturday': '周六', 'sunday': '周日'}
PART = {'morning': '上午', 'afternoon': '下午', 'evening': '晚上', 'night': '晚上'}

# ---- 交通方式归一化: 原始表里同一含义有至少五种写法 ----
TRANSPORT_RULES = [
    (r'\btaxi\b|打车|出租', '打车'),
    (r'\bcar\b|自驾|开车', '自驾'),
    (r'\bbus\b|公交|公车', '公交'),
    (r'\bbike\b|bicycl|自行车|骑行|单车', '自行车'),
]
FLEXIBLE = {'自驾', '打车'}   # 视为可以接受更远通勤

# ---- 英文站名 / 简写的人工别名. 无法自动解析的才写进来 ----
STATION_ALIASES = {
    'shanghaiuniversityoffinanceandeconomicsstation': '上海财经大学',
    'fengpuavenue': '奉浦大道',
    'shiguangroad': '市光路',
}


def norm_slot(text):
    """'Friday Evening (After 6pm)' / 'Saturday morning' -> '周五晚上' / '周六上午'"""
    t = str(text).lower()
    t = re.sub(r'[（(].*?[)）]', ' ', t)
    d = next((v for k, v in DAY.items() if k in t), None)
    p = next((v for k, v in PART.items() if k in t), None)
    if d and p:
        return d + p
    return None


def norm_transport(text):
    t = str(text or '').lower()
    out = []
    for pat, val in TRANSPORT_RULES:
        if re.search(pat, t) and val not in out:
            out.append(val)
    return out or ['无']


LINE_PAT = re.compile(r'(?:line\s*|(\d+)\s*号线)|line\s*(\d+)', re.I)


def split_station_tokens(text):
    """把一个单元格拆成若干「线路号 + 站名」候选片段.

    需要应付的写法(全部来自原始表):
      Line 10‐殷高东路 Yin'ao East Road
      Line 10 国权路 Guoquan Road \n and \n Line 18 Shanghai university ...
      Line 15 永德路 Yongde Lu \n or \n Line 8 沈杜公路 Shendu road
      11号线/15号线 上海西站
      13号线 金运路/ 2号线 虹桥火车站
      15号线永德路或8号线沈杜公路
      3号线 东宝兴路或8号线西藏北路
      11号线陈翔公路站
    """
    s = str(text or '')
    s = s.replace('‐', '-').replace('－', '-')
    # 统一分隔符. 注意不能全局按 and / or 切分, 否则
    # 「Shanghai university of finance and economics station」会被切成两半
    s = re.sub(r'\n+', '|', s)
    s = s.replace('或', '|').replace('/', '|').replace('、', '|').replace(';', '|')
    parts = []
    for p in s.split('|'):
        p = p.strip(' -|')
        if not p or re.fullmatch(r'(and|or)', p, flags=re.I):
            continue
        n_line = len(re.findall(r'\d+\s*号线|line\s*\d+', p, flags=re.I))
        if n_line >= 2:   # 同一片段里出现多个线路号才按 and/or 再切
            parts += [q.strip(' -') for q in re.split(r'\b(?:and|or)\b', p, flags=re.I) if q.strip(' -')]
        else:
            parts.append(p)

    # 「11号线|15号线 上海西站」这类: 前面的片段只有线路号没有站名,
    # 需要把后面片段的站名借过来
    toks = []
    pending_lines = []
    for p in parts:
        lines = re.findall(r'(\d+)\s*号线', p) + re.findall(r'line\s*(\d+)', p, flags=re.I)
        body = re.sub(r'\d+\s*号线|line\s*\d+', ' ', p, flags=re.I).strip(' -')
        has_name = bool(re.search(r'[一-鿿]', body)) or bool(re.search(r'[a-z]{4}', body, re.I))
        if not has_name and lines:
            pending_lines += lines
            continue
        for ln in (lines or [None]):
            toks.append((ln, body))
        for ln in pending_lines:
            toks.append((ln, body))
        pending_lines = []
    return toks


# 匹配方式的严重程度: exact/zh-substr 属于这种「中文站名 + 英文音译」格式的正常情况
SEVERITY = {'exact': None, 'zh-substr': None,
            'alias': 'WARN', 'pinyin': 'WARN', 'line-scan': 'WARN'}


def resolve_stations(engine, text, warn, ctx):
    """一个单元格 -> 去重后的站名列表"""
    out = []
    for ln, body in split_station_tokens(text):
        st, how = engine.resolve(body, line_no=ln, aliases=STATION_ALIASES)
        if st is None:
            warn.append(('ERROR', f"{ctx}: 站名无法解析 -> {body!r} (线路 {ln})"))
            continue
        if st not in out:
            out.append(st)
        sev = SEVERITY.get(how)
        if sev:
            warn.append((sev, f"{ctx}: 站名靠 {how} 推断 -> {body!r} => {st}, 请人工确认"))
        if ln and ln not in engine.line_of_station[st]:
            warn.append(('ERROR', f"{ctx}: 线路与站点不一致 -> 填 {ln} 号线, 但 {st} 属于 "
                                  f"{sorted(engine.line_of_station[st])} 号线"))
    return out


def load(path, engine):
    wb = openpyxl.load_workbook(path, data_only=True)
    warn = []

    # ---------------- 志愿者 ----------------
    vols = []
    ws = wb['vol info']
    header = [str(c.value or '').lower() for c in ws[1]]
    old_transport_col = len(header) > 2 and 'transport' in header[2]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        vid = str(row[0]).strip()
        ctx = f"志愿者 {vid}"
        extra_station_raw = '' if old_transport_col else row[2]
        station_raw = '; '.join(str(x) for x in (row[1], extra_station_raw)
                                if str(x or '').strip())
        stations = resolve_stations(engine, station_raw, warn, ctx)
        transport = ['无']
        slots, bad = [], []
        for piece in str(row[3] or '').split('┋'):
            piece = piece.strip()
            if not piece:
                continue
            s = norm_slot(piece)
            (slots if s else bad).append(s or piece)
        for b in bad:
            warn.append(('ERROR', f"{ctx}: 时段无法解析 -> {b!r}"))
        if not stations:
            warn.append(('ERROR', f"{ctx}: 没有可用站点, 该志愿者无法参与求解"))
        vols.append({
            'id': vid, 'stations': stations, 'transport': transport,
            'flexible': False,
            'slots': sorted(set(slots)),
            'raw_station': station_raw, 'raw_transport': '',
        })

    # ---------------- 项目点与班级 ----------------
    classes = []
    sites = {}
    ws = wb['site info']
    cur_sid = None
    cur_station_raw = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ''):
            cur_sid = str(row[0]).strip()
            cur_station_raw = row[1]
            warn_ctx = f"项目点 {cur_sid}"
            sites[cur_sid] = {
                'id': cur_sid,
                'raw_station': str(row[1]).replace('\n', ' '),
                'stations': resolve_stations(engine, row[1], warn, warn_ctx),
            }
        if row[2] in (None, ''):
            continue
        if cur_sid is None:
            warn.append(('ERROR', f"班级 {row[2]} 找不到所属项目点, 已跳过"))
            continue
        slot = norm_slot(row[3])
        if slot is None:
            warn.append(('ERROR', f"项目点 {cur_sid} {row[2]}: 上课时段无法解析 -> {row[3]!r}"))
            continue
        classes.append({
            'site': cur_sid, 'cls': str(row[2]).strip(), 'slot': slot,
            'need': int(row[4] or 2),
            'stations': sites[cur_sid]['stations'],
            'raw_slot': row[3],
        })

    # 项目点编号断档检查
    nums = sorted(int(s) for s in sites if str(s).isdigit())
    missing = [n for n in range(min(nums), max(nums) + 1) if n not in nums]
    if missing:
        warn.append(('WARN', f"项目点编号断档: 缺 {missing}, 无法从数据判断是录入遗漏还是项目点已关闭"))

    return vols, classes, sites, warn


if __name__ == '__main__':
    import sys
    from metro import TravelTimeEngine
    eng = TravelTimeEngine('sh_metro_raw.json')
    vols, classes, sites, warn = load(sys.argv[1], eng)
    print(f"志愿者 {len(vols)}  项目点 {len(sites)}  班级 {len(classes)}\n")
    print("---- 清洗告警 ----")
    for lv in ('ERROR', 'WARN'):
        hits = [m for l, m in warn if l == lv]
        print(f"  [{lv}] {len(hits)} 条")
        for m in hits:
            print('    ', m)
    print()
    print("---- 志愿者样例 ----")
    for v in vols[:3] + [v for v in vols if v['id'] in ('Vol 37', 'Vol 38')]:
        print(f"  {v['id']:8} {v['stations']}  {v['transport']} flex={v['flexible']}  {v['slots']}")
    print("\n---- 项目点解析 ----")
    for s in sites.values():
        print(f"  {s['id']:>3}  {str(s['stations']):40} <= {s['raw_station'][:52]}")
